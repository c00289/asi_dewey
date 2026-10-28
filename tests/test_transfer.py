from __future__ import annotations

import json

import pytest
from openpyxl import load_workbook

from asi_dewey import Brain
from tests.conftest import WORKBOOK


def test_json_dump_and_load_modes(seeded, tmp_path):
    out = tmp_path / "dump.json"
    dumped = seeded.execute("dump", out=str(out))
    assert dumped.ok and dumped.data["entities"]["fact"] == 2
    payload = json.loads(out.read_text(encoding="utf-8"))
    assert payload["schema_version"] == 1 and set(payload["entities"]) >= {"fact", "task", "safety_plan_row"}

    with Brain(tmp_path / "other.db") as other:
        loaded = other.execute("load", path=str(out), mode="replace")
        assert loaded.ok, loaded.errors
        assert other.execute("count", entity="fact").data["count"] == 2
        # merge keeps ids, append makes new ones
        assert other.execute("load", path=str(out), mode="merge").data["entities"]["fact"]["updated"] == 2
        assert other.execute("count", entity="fact").data["count"] == 2
        other.execute("load", path=str(out), mode="append")
        assert other.execute("count", entity="fact").data["count"] == 4

    single = tmp_path / "facts.json"
    assert seeded.execute("dump", entity="fact", out=str(single)).ok
    assert json.loads(single.read_text(encoding="utf-8"))["entity"] == "fact"


def test_excel_export_and_reimport(seeded, tmp_path):
    exported = seeded.execute("export", format="excel")
    assert exported.ok, exported.errors
    path = exported.data["path"]
    wb = load_workbook(path)
    assert set(wb.sheetnames) == {
        "tasks", "facts", "WL functions", "WL use cases", "links", "JIRA", "QMS", "people", "Safety Plan",
    }
    ws = wb["Safety Plan"]
    headers = [c.value for c in ws[1]]
    assert headers[:4] == ["order", "PDP Phases", "Safety Lifecycle Phase", "Safety Deliverables"]
    assert headers.count("Tech Lead - Functional Safety") == 2 and len(headers) == 63
    assert ws.freeze_panes == "A2"
    assert wb["links"]["D2"].hyperlink.target == "https://x.y/z"
    assert wb["tasks"]["D2"].value.date().isoformat() == "2026-10-20"

    with Brain(tmp_path / "re.db") as other:
        result = other.execute("import", format="excel", path=path, mode="replace")
        assert result.ok, result.errors
        row = other.execute("list", entity="safety_plan_row").data[0]
        assert row["deliverable"] == "Safety Plan" and row["loc_dms"] and row["reviewer_tech_lead_functional_safety"]
        assert row["author_tech_lead_functional_safety"] is False
        assert other.execute("list", entity="task").data[0]["start"] == "2026-10-20"


def test_locked_workbook_falls_back_to_sibling(seeded, tmp_path, monkeypatch):
    import asi_dewey.io.exporters as ex

    calls = []

    def fake_save(path):
        calls.append(path)
        if len(calls) == 1:
            raise PermissionError("locked")
        path.write_bytes(b"ok")

    path, warning = ex._save_with_fallback(fake_save, tmp_path / "view.xlsx", ".xlsx")
    assert path != tmp_path / "view.xlsx" and path.suffix == ".xlsx" and "locked" in warning


def test_run_script(seeded, tmp_path):
    script = tmp_path / "steps.jsonl"
    script.write_text(
        '# comment\n'
        '{"verb": "create", "args": {"entity": "fact", "data": {"name": "from script"}}}\n'
        '{"verb": "count", "args": {"entity": "fact"}}\n'
        '{"verb": "get", "args": {"entity": "fact", "id": 999}}\n',
        encoding="utf-8",
    )
    result = seeded.execute("run", path=str(script))
    assert not result.ok and result.data["steps"] == 3 and result.data["failed"] == [3]
    assert result.data["results"][1]["data"]["count"] == 3
    assert seeded.storage.audit_tail(3)[-1]["actor"] == "python"
    stop = seeded.execute("run", path=str(script), stop_on_error=True)
    assert stop.data["steps"] == 3


@pytest.mark.skipif(not WORKBOOK.is_file(), reason="real workbook not present")
def test_seed_from_real_workbook(tmp_path):
    with Brain(tmp_path / "seed.db") as b:
        result = b.execute("import", format="excel", path=str(WORKBOOK), mode="replace")
        assert result.ok, result.errors
        counts = b.execute("status").data["entities"]
        assert counts["task"] >= 10 and counts["safety_plan_row"] >= 48 and counts["machine_function"] >= 30
        plan = b.execute("list", entity="safety_plan_row", filter=["order=1"]).data[0]
        assert plan["lifecycle_phase"] == "1. Concept Development" and plan["ref_qf_17"] is True
        assert b.execute("list", entity="task", filter=["order=1"]).data[0]["start"] == "2026-10-19"
        out = b.execute("export", format="excel", out=str(tmp_path / "regen.xlsx"))
        assert out.ok and out.data["entities"]["safety_plan_row"] == counts["safety_plan_row"]
