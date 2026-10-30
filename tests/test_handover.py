"""The handover (0.11.0): a schema-4 database's STPA work written out in the shapes asi_kletz 0.24.0 reads.

The source is ``tests/fixtures/dewey-0.10.1.db``, a synthetic 0.10.1 database (schema 4) built from the
0.9.0 fixture with 0.10.1 itself: three control structures (the eight-box synthetic one, an empty one, the
three-box gate), one STPA analysis read from the synthetic conveyor-gate page and linked to the gate
structure, one note typed on UCA1. Every name is invented. The last test runs the sister tool itself when
its checkout is beside this one (``../asi_kletz/.venv``) and is skipped elsewhere.
"""

from __future__ import annotations

import json
import os
import shutil
import sqlite3
import subprocess
from pathlib import Path

import pytest
from openpyxl import load_workbook

from asi_dewey import Brain, __version__
from asi_dewey.automation.cli import main
from asi_dewey.handover import (
    CS_ENTITIES,
    ENTITIES,
    JSON_FILE,
    MANIFEST_FILE,
    STPA_ENTITIES,
    TABLES,
    WORKBOOK_FILE,
    decode_list,
    read_tables,
)

from .conftest import REPO_ROOT

FIXTURE_0_10_1 = REPO_ROOT / "tests" / "fixtures" / "dewey-0.10.1.db"
KLETZ = REPO_ROOT.parent / "asi_kletz" / ".venv" / "Scripts" / "asi-kletz.exe"
COUNTS = {
    "control_structure": 3, "cs_controller": 11, "cs_algorithm": 3, "cs_process_model_entry": 2, "cs_edge": 11,
    "cs_signal": 13, "cs_note": 1,
    "stpa_analysis": 1, "stpa_loss": 2, "stpa_hazard": 5, "stpa_uca": 8, "stpa_scenario": 10, "stpa_constraint": 6,
}
SYNTHETIC = {"controllers": 8, "algorithms": 3, "process_model_entries": 2, "edges": 9, "signals": 10, "notes": 1}
BUNDLE = Path("evidence") / "conveyor-gate-stpa-2026-10-29"  # as the fixture's analysis records it (relative)


@pytest.fixture
def schema4(tmp_path):
    db = tmp_path / "dewey.db"
    shutil.copy(FIXTURE_0_10_1, db)
    return db


def _table_counts(db: Path) -> dict[str, int]:
    con = sqlite3.connect(db)
    try:
        return {t.entity: con.execute(f'SELECT count(*) FROM "{t.table}"').fetchone()[0] for t in TABLES}
    finally:
        con.close()


def write_stub_bundle(folder: Path) -> Path:
    """The three files a Hopper bundle has, with nothing real in them: enough for the copy to be checked."""
    (folder / "pages").mkdir(parents=True, exist_ok=True)
    manifest = {"kind": "confluence", "pages": [{"id": "900000001"}]}
    (folder / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
    (folder / "pages.json").write_text("[]", encoding="utf-8")
    (folder / "pages" / "900000001.md").write_text("# Conveyor Gate Safety Function STPA\n", encoding="utf-8")
    return folder


# -- the frozen shape ----------------------------------------------------------------------------------
def test_the_frozen_shape_is_the_thirteen_of_0_10_1():
    assert ENTITIES == (*CS_ENTITIES, *STPA_ENTITIES) and len(ENTITIES) == 13
    assert [t.sheet for t in TABLES] == [
        "CS structures", "CS controllers", "CS algorithms", "CS process model", "CS edges", "CS signals", "CS notes",
        "STPA analyses", "STPA losses", "STPA hazards", "STPA UCAs", "STPA scenarios", "STPA constraints",
    ]
    uca = next(t for t in TABLES if t.entity == "stpa_uca")
    assert uca.headers[:3] == ["analysis", "order", "code"] and uca.headers[-1] == "uid"
    assert {c.name for c in uca.columns if c.kind == "json"} == {"hazards"}
    assert decode_list(None) == [] and decode_list('["H2", "H2.1"]') == ["H2", "H2.1"] and decode_list("H2") == ["H2"]


def test_read_tables_decodes_what_0_10_1_stored(schema4):
    assert _table_counts(schema4) == COUNTS
    with Brain(schema4) as b:
        rows = read_tables(b)
    assert {e: len(r) for e, r in rows.items()} == COUNTS
    uca1 = next(u for u in rows["stpa_uca"] if u["code"] == "UCA1")
    assert uca1["hazards"] == ["H2", "H2.1"] and uca1["retired"] is False and "id" not in uca1
    assert uca1["cs_controller"] == "gate_controller" and uca1["cs_signal"] == "CA-10"
    assert uca1["note"] == "explored in Dewey; to be sanitised in Kletz" and uca1["uid"] and uca1["created_at"]
    assert [s["slug"] for s in rows["control_structure"]] == ["gate", "demo_loader_control_structure", "demo_structure"]
    placeholder = next(s for s in rows["cs_signal"] if s["signal_id"] == "CA-X1")
    assert placeholder["placeholder"] is True
    entry = rows["cs_process_model_entry"][0]
    assert entry["control_actions"] == ["CA-01"] and entry["updated_by"] == ["FB-01"]


# -- the verb -------------------------------------------------------------------------------------------
def test_handover_writes_the_files_kletz_reads(schema4, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)  # the fixture's bundle path is relative and not here: a warning, not a failure
    out = tmp_path / "hand"
    with Brain(schema4) as b:
        result = b.execute("handover", out=str(out), actor="owner")
        assert result.ok, result.errors
        data = result.data
        assert data["counts"] == COUNTS and data["out"] == str(out) and data["database"] == str(schema4)
        assert result.warnings == [
            f"analysis conveyor_gate: its bundle {BUNDLE} is not on this machine; Kletz loads the rows from "
            f"{JSON_FILE} and cannot read the page again until the bundle is beside it"
        ]
        audit = b.storage.audit_tail(1)[0]
        assert audit["action"] == "handover" and audit["actor"] == "owner" and audit["payload"]["stpa_uca"] == 8
        assert b.events.last(1)[0]["topic"] == "handover.written"
        # the thirteen tables are still there, untouched: nothing is deleted by a handover
        assert _table_counts(schema4) == COUNTS

    # the dump envelope: the thirteen entities in parent-first order, uids kept, ids dropped
    payload = json.loads((out / JSON_FILE).read_text(encoding="utf-8"))
    assert list(payload["entities"]) == list(ENTITIES) and payload["app_version"] == __version__
    assert payload["handover"]["for"].startswith("asi_kletz 0.24.0 load")
    for entity, rows in payload["entities"].items():
        assert len(rows) == COUNTS[entity]
        assert all("id" not in r and r["uid"] for r in rows), entity
    uca1 = next(u for u in payload["entities"]["stpa_uca"] if u["code"] == "UCA1")
    assert uca1["hazards"] == ["H2", "H2.1"] and uca1["retired"] is False and uca1["cs_signal"] == "CA-10"
    assert payload["entities"]["stpa_analysis"][0]["control_structure"] == "gate"

    # one datasets folder per structure in the Roy's Brain envelope, the script's names on the rows
    assert [s["slug"] for s in data["structures"]] == ["gate", "demo_loader_control_structure", "demo_structure"]
    demo = out / "control-structures" / "demo_loader_control_structure"
    files = {name: json.loads((demo / f"{name}.json").read_text(encoding="utf-8")) for name in SYNTHETIC}
    assert {name: f["row_count"] for name, f in files.items()} == SYNTHETIC
    assert all(f["entity"] == name and len(f["rows"]) == f["row_count"] for name, f in files.items())
    assert files["controllers"]["rows"][0] == {
        "id": "remote_operator", "order": 1, "name": "Remote Operator", "kind": "human_controller",
        "responsibility": "Starts and stops the mission.\nWatches the site.",
        "process_model_title": "Operator process model", "note": "",
    }
    o1 = files["algorithms"]["rows"][0]
    assert o1["controller_id"] == "remote_operator" and o1["id"] == "O1"
    e1 = next(e for e in files["edges"]["rows"] if e["id"] == "e1")
    assert e1["source_id"] == "remote_operator" and e1["target_id"] == "vehicle_controller"
    assert e1["signal_ids"] == ["CA-01", "CA-02"]
    x1 = next(s for s in files["signals"]["rows"] if s["signal_id"] == "CA-X1")
    assert x1["edge_id"] == "e8" and x1["placeholder"] is True and x1["kind"] == "control_action"
    assert files["process_model_entries"]["rows"][0]["controller_id"] == "remote_operator"
    assert files["notes"]["rows"][0]["attached_to"] == "vehicle_controller"
    head = json.loads((demo / "structure.json").read_text(encoding="utf-8"))
    assert head["name"] == "Demo Loader Control Structure" and head["system"] == "Demo loader" and head["uid"]
    empty = json.loads((out / "control-structures" / "demo_structure" / "controllers.json").read_text(encoding="utf-8"))
    assert empty["row_count"] == 0 and empty["rows"] == []  # an empty structure still gets the two required files
    assert (out / "control-structures" / "demo_structure" / "edges.json").is_file()

    # the workbook: the thirteen sheets, headed as the 0.10.1 workbook was, the uid column last and hidden
    wb = load_workbook(out / WORKBOOK_FILE)
    assert wb.sheetnames == [t.sheet for t in TABLES]
    ws = wb["STPA UCAs"]
    headers = [c.value for c in ws[1]]
    assert headers[:4] == ["analysis", "order", "code", "controller"] and headers[-1] == "uid"
    assert ws.column_dimensions["O"].hidden and ws.max_row == 9
    assert ws["H2"].value == '["H2", "H2.1"]' and ws["I2"].value == 0  # hazards as JSON text, retired as 0

    # the manifest and the Kletz command lines
    manifest = json.loads((out / MANIFEST_FILE).read_text(encoding="utf-8"))
    assert manifest["counts"] == COUNTS and manifest["kletz"] == data["kletz"]
    assert manifest["warnings"] == result.warnings
    assert data["kletz"][0] == f'asi-kletz load "{out / JSON_FILE}" --mode merge'
    assert data["kletz"][2] == (
        f'asi-kletz cs new from_datasets --from "{demo}" --name "Demo Loader Control Structure" '
        "--slug demo_loader_control_structure --system \"Demo loader\""
    )
    assert len(data["kletz"]) == 4 and data["analyses"][0]["bundle"] is None  # no bundle here: no from_bundle line
    assert data["analyses"][0]["counts"] == {
        "stpa_loss": 2, "stpa_hazard": 5, "stpa_uca": 8, "stpa_scenario": 10, "stpa_constraint": 6,
    }


def test_handover_copies_the_bundle_when_it_is_on_this_machine(schema4, tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    write_stub_bundle(tmp_path / BUNDLE)
    with Brain(schema4) as b:
        result = b.execute("handover", out=str(tmp_path / "hand"))
    assert result.ok and result.warnings == []
    copied = tmp_path / "hand" / "bundles" / "conveyor_gate"
    assert (copied / "manifest.json").is_file() and (copied / "pages" / "900000001.md").is_file()
    assert result.data["analyses"][0]["bundle"] == str(copied)
    assert result.data["kletz"][-1] == (
        f'asi-kletz stpa new from_bundle --from "{copied}" --name "Conveyor Gate" --slug conveyor_gate'
    )
    # a bundle beside the database counts too (the working directory elsewhere), and a second run replaces the copy
    (tmp_path / "elsewhere").mkdir()
    monkeypatch.chdir(tmp_path / "elsewhere")
    with Brain(schema4) as b:
        again = b.execute("handover", out=str(tmp_path / "hand"))
    assert again.ok and again.warnings == [] and (copied / "manifest.json").is_file()


def test_handover_default_folder_and_nothing_to_hand_over(schema4, tmp_path):
    with Brain(tmp_path / "fresh.db") as fresh:  # a 0.11.0 database: no such tables
        refused = fresh.execute("handover", out=str(tmp_path / "nothing"))
    assert not refused.ok and refused.errors[0].startswith("nothing to hand over")
    assert not (tmp_path / "nothing").exists()
    with Brain(schema4, out_dir=tmp_path / "out") as b:
        result = b.execute("handover")
    assert result.ok and result.data["out"] == str(tmp_path / "out" / "handover")
    assert (tmp_path / "out" / "handover" / JSON_FILE).is_file()


def test_handover_through_the_cli_and_the_schema(schema4, tmp_path, capsys):
    assert main(["--db", str(schema4), "--no-forward", "handover", "--out", str(tmp_path / "cli")]) == 0
    printed = json.loads(capsys.readouterr().out)
    assert printed["ok"] and printed["data"]["counts"] == COUNTS and printed["verb"] == "handover"
    with Brain(schema4) as b:
        verbs = {c["verb"]: c for c in b.execute("schema").data["commands"]}
        assert "handover" in verbs and "cs" not in verbs and "stpa" not in verbs
        by_name = {a["name"]: a for a in verbs["handover"]["arguments"]}
        assert list(by_name) == ["out"] and by_name["out"]["type"] == "path" and by_name["out"]["flag"] is True


# -- the sister tool itself ----------------------------------------------------------------------------
def _kletz(env: dict[str, str], *args: str) -> dict:
    proc = subprocess.run(
        [str(KLETZ), "--no-forward", *args], capture_output=True, text=True, env=env, timeout=120, check=False,
    )
    assert proc.returncode in (0, 1), proc.stderr
    return json.loads(proc.stdout)


@pytest.mark.skipif(not KLETZ.is_file(), reason="asi_kletz is not checked out beside this repo")
def test_kletz_reads_the_handover_with_the_same_counts(schema4, tmp_path, monkeypatch):
    """The round trip the family rule needs: ``load`` takes the rows (counts equal, the note, the link and the
    uids kept); ``cs new from_datasets`` builds the synthetic structure through Kletz's own factory."""
    monkeypatch.chdir(tmp_path)
    out = tmp_path / "hand"
    with Brain(schema4) as b:  # closed before Kletz runs: one primary tool at a time
        data = b.execute("handover", out=str(out)).data
    env = {
        **os.environ,  # the per-test lock directory of conftest, so the sibling never touches the real one
        "ASI_KLETZ_CONFIG": str(tmp_path / "absent.toml"), "ASI_KLETZ_DB": str(tmp_path / "kletz.db"),
        "ASI_KLETZ_WORKBOOK": str(tmp_path / "kletz.xlsx"), "ASI_KLETZ_OUT": str(tmp_path / "kletz-out"),
        "ASI_KLETZ_CONTROL_STRUCTURES": str(tmp_path / "no-cs"), "ASI_KLETZ_STANDARDS": str(tmp_path / "no-standards"),
    }
    loaded = _kletz(env, "load", str(out / JSON_FILE), "--mode", "merge")
    assert loaded["ok"], loaded["errors"]
    assert {e: s["inserted"] for e, s in loaded["data"]["entities"].items() if e in COUNTS} == COUNTS
    status = _kletz(env, "status")
    assert {e: n for e, n in status["data"]["entities"].items() if e in COUNTS} == COUNTS
    got = _kletz(env, "stpa", "get", "conveyor_gate")
    assert got["ok"] and got["data"]["control_structure"] == "gate"
    uca1 = next(u for u in got["data"]["parts"]["ucas"] if u["code"] == "UCA1")
    written = json.loads((out / JSON_FILE).read_text(encoding="utf-8"))["entities"]["stpa_uca"]
    expected = next(u for u in written if u["code"] == "UCA1")
    assert (uca1["note"], uca1["cs_controller"], uca1["cs_signal"], uca1["hazards"], uca1["uid"]) == (
        expected["note"], "gate_controller", "CA-10", ["H2", "H2.1"], expected["uid"],
    )
    # a second load updates by uid: nothing duplicated
    again = _kletz(env, "load", str(out / JSON_FILE), "--mode", "merge")
    assert again["ok"] and all(s["inserted"] == 0 for e, s in again["data"]["entities"].items() if e in COUNTS)
    assert {e: n for e, n in _kletz(env, "status")["data"]["entities"].items() if e in COUNTS} == COUNTS
    # the datasets route, through Kletz's factory, into a second database
    env["ASI_KLETZ_DB"] = str(tmp_path / "kletz2.db")
    demo = next(s for s in data["structures"] if s["slug"] == "demo_loader_control_structure")
    built = _kletz(env, "cs", "new", "from_datasets", "--from", demo["folder"], "--name", demo["name"],
                   "--slug", demo["slug"], "--system", "Demo loader")
    assert built["ok"], built["errors"]
    assert built["data"]["counts"] == SYNTHETIC and built["data"]["slug"] == "demo_loader_control_structure"
    assert built["warnings"] == [
        "cs_edge 8: edge e8: control action from fleet_server points up the hierarchy to vehicle_controller"
    ]
