"""Handover: the STPA work explored here, written out in the shapes asi_kletz 0.24.0 reads.

The owner, 2026-10-29: "control structures were supposed to go into Kletz and not Dewey"; asked where STPA
should live, "move everything to Kletz"; and the family rule: "Dewey is more for experimental and day to
day; all true work products will end up in Kletz once sanitized." So 0.11.0 drops the seven
control-structure entities of 0.9.0 and the six STPA entities of 0.10.0 from the registry and keeps the
one thing the rule needs: a way to hand what was explored here over to Kletz. The thirteen tables stay in
an older database, unregistered and untouched (no table is dropped, no row is lost for anyone who has not
handed over yet); ``write_handover`` reads them with plain SQL, through the frozen schema-4 shape in
``TABLES`` (the entity names, tables, sheets and fields exactly as 0.10.1 declared them), and writes:

- ``dewey-stpa.json``, the dump envelope (``{"entities": {<entity>: [rows]}}``, the thirteen entities in
  parent-first order, every row with its ``uid`` and without its ``id``) that ``asi-kletz load --mode
  merge`` reads: the rows arrive as they were explored, notes and links included, a row already handed
  over is updated by its uid, and no id can collide with a row Kletz already has;
- ``control-structures/<slug>/`` per structure, the six JSON files of the Roy's Brain envelope
  (``controllers.json`` ... ``notes.json``, each ``{"entity", "row_count", "rows"}`` with the script's
  names: ``id``, ``controller_id``, ``edge_id``, ``source_id``, ``target_id``) that ``asi-kletz cs new
  from_datasets --from <folder>`` reads through the factory, where a duplicate signal id is a warning;
- ``bundles/<slug>/`` per analysis whose Hopper evidence bundle is still on this machine, a copy, so
  ``asi-kletz stpa new from_bundle --from <folder>`` can read the page again there;
- ``dewey-stpa.xlsx``, the thirteen sheets headed as the workbook was (the uid column last and hidden),
  for a person to look at and for ``asi-kletz import excel``;
- ``handover.json``, the manifest: what was written, the counts, and the Kletz command lines.

Nothing is deleted here. The whole handover is audited once as ``handover`` and published as
``handover.written``. A database that never held the tables, or holds them empty, has nothing to hand over
and the verb says so.
"""

from __future__ import annotations

import datetime as dt
import json
import shutil
from dataclasses import dataclass
from pathlib import Path
from typing import TYPE_CHECKING, Any

from .errors import CommandError
from .version import SCHEMA_VERSION, __version__

if TYPE_CHECKING:
    from .brain import Brain

KLETZ_VERSION = "0.24.0"  # the asi_kletz release whose load, from_datasets and from_bundle read these files
JSON_FILE = "dewey-stpa.json"
WORKBOOK_FILE = "dewey-stpa.xlsx"
MANIFEST_FILE = "handover.json"
STRUCTURES_DIR = "control-structures"
BUNDLES_DIR = "bundles"
DEFAULT_DIR = "handover"  # under Brain.out_dir when --out is not given
DATASET_FILES = ("controllers", "algorithms", "process_model_entries", "signals", "edges", "notes")


@dataclass(frozen=True)
class Column:
    name: str
    header: str
    kind: str = "text"  # text | int | bool | json


@dataclass(frozen=True)
class Table:
    """One of the thirteen entities as 0.10.1 declared it (schema 4): the frozen shape the handover reads."""

    entity: str
    table: str
    sheet: str
    order_by: tuple[str, ...]
    columns: tuple[Column, ...]

    @property
    def headers(self) -> list[str]:
        return [*(c.header for c in self.columns), "uid"]


def _c(name: str, header: str | None = None, kind: str = "text") -> Column:
    return Column(name, header or name, kind)


_STRUCTURE = _c("structure")
_ANALYSIS = _c("analysis")
_ORDER = _c("order", kind="int")
_CITE = _c("cite")
_NOTE = _c("note")

TABLES: tuple[Table, ...] = (
    Table("control_structure", "control_structures", "CS structures", ("name", "id"), (
        _c("name"), _c("slug"), _c("system"), _c("description"), _c("source"), _c("status"),
        _c("revision", kind="int"),
    )),
    Table("cs_controller", "cs_controllers", "CS controllers", ("structure", "order", "id"), (
        _STRUCTURE, _ORDER, _c("key"), _c("name"), _c("kind"), _c("responsibility"),
        _c("process_model_title", "process model title"), _NOTE,
    )),
    Table("cs_algorithm", "cs_algorithms", "CS algorithms", ("structure", "controller", "order", "id"), (
        _STRUCTURE, _c("controller"), _ORDER, _c("code"), _c("name"),
    )),
    Table("cs_process_model_entry", "cs_process_model_entries", "CS process model",
          ("structure", "controller", "order", "id"), (
        _STRUCTURE, _c("controller"), _ORDER, _c("heading"), _c("control_actions", "control actions", "json"),
        _c("variables", kind="json"), _c("updated_by", "updated by", "json"), _c("references", kind="json"),
    )),
    Table("cs_edge", "cs_edges", "CS edges", ("structure", "order", "id"), (
        _STRUCTURE, _ORDER, _c("key"), _c("kind"), _c("source"), _c("target"),
    )),
    Table("cs_signal", "cs_signals", "CS signals", ("structure", "edge", "order", "id"), (
        _STRUCTURE, _c("edge"), _ORDER, _c("signal_id", "signal id"), _c("kind"), _c("name"),
        _c("placeholder", kind="bool"),
    )),
    Table("cs_note", "cs_notes", "CS notes", ("structure", "order", "id"), (
        _STRUCTURE, _ORDER, _c("text"), _c("attached_to", "attached to"),
    )),
    Table("stpa_analysis", "stpa_analyses", "STPA analyses", ("name", "id"), (
        _c("name"), _c("slug"), _c("system"), _c("status"), _c("control_structure", "control structure"),
        _c("source_kind", "source kind"), _c("source_page_id", "source page id"), _c("source_url", "source url"),
        _c("source_version", "source version", "int"), _c("source_body_sha256", "source body sha256"),
        _c("source_taken_at", "source taken at"), _c("source_bundle", "source bundle"), _NOTE,
    )),
    Table("stpa_loss", "stpa_losses", "STPA losses", ("analysis", "order", "id"), (
        _ANALYSIS, _ORDER, _c("code"), _c("text"), _CITE, _NOTE,
    )),
    Table("stpa_hazard", "stpa_hazards", "STPA hazards", ("analysis", "order", "id"), (
        _ANALYSIS, _ORDER, _c("code"), _c("text"), _c("parent"), _c("losses", kind="json"),
        _c("retired", kind="bool"), _CITE, _NOTE,
    )),
    Table("stpa_uca", "stpa_ucas", "STPA UCAs", ("analysis", "order", "id"), (
        _ANALYSIS, _ORDER, _c("code"), _c("controller"), _c("control_action", "control action"), _c("category"),
        _c("text"), _c("hazards", kind="json"), _c("retired", kind="bool"), _c("duplicate_of", "duplicate of"),
        _c("cs_controller", "cs controller"), _c("cs_signal", "cs signal"), _CITE, _NOTE,
    )),
    Table("stpa_scenario", "stpa_scenarios", "STPA scenarios", ("analysis", "order", "id"), (
        _ANALYSIS, _ORDER, _c("row", kind="int"), _c("uca"), _c("counterpart"), _c("types", kind="json"),
        _c("causal_factors", "causal factors"), _c("beliefs"), _c("constraints", kind="json"),
        _c("retired", kind="bool"), _CITE, _NOTE,
    )),
    Table("stpa_constraint", "stpa_constraints", "STPA constraints", ("analysis", "order", "id"), (
        _ANALYSIS, _ORDER, _c("code"), _c("text"), _CITE, _NOTE,
    )),
)
ENTITIES = tuple(t.entity for t in TABLES)
CS_ENTITIES = ENTITIES[:7]
STPA_ENTITIES = ENTITIES[7:]
_PARTS = {  # entity -> the dataset key its rows go under
    "cs_controller": "controllers", "cs_algorithm": "algorithms", "cs_process_model_entry": "process_model_entries",
    "cs_edge": "edges", "cs_signal": "signals", "cs_note": "notes",
}


def _timestamp() -> str:
    return dt.datetime.now(dt.UTC).replace(microsecond=0).isoformat()


def decode_list(text: Any) -> list[Any]:
    """The list behind a JSON column's text: NULL or blank is ``[]``; text that is not a JSON array (never
    written by this tool) is kept as a one-item list rather than lost."""
    if text is None or text == "":
        return []
    if isinstance(text, list):
        return text
    try:
        value = json.loads(text)
    except (TypeError, ValueError):
        return [str(text)]
    return value if isinstance(value, list) else [value]


def _decode(table: Table, raw: dict[str, Any]) -> dict[str, Any]:
    """A stored row as the API answered it in 0.10.1: lists as lists, bools as bools, no ``id``."""
    row: dict[str, Any] = {}
    for column in table.columns:
        value = raw.get(column.name)
        if column.kind == "json":
            value = decode_list(value)
        elif column.kind == "bool":
            value = bool(value)
        row[column.name] = value
    row["uid"] = raw.get("uid")
    for key in ("created_at", "updated_at"):
        if raw.get(key) is not None:
            row[key] = raw[key]
    return row


def _quote(name: str) -> str:
    return '"' + name.replace('"', '""') + '"'


def read_tables(brain: Brain) -> dict[str, list[dict[str, Any]]]:
    """The rows of every one of the thirteen tables still in the database, decoded, parents first; a
    table that is not there (a database that never ran 0.9.0 to 0.10.1) reads as no rows."""
    present = {
        r["name"] for r in brain.storage.fetchall("SELECT name FROM sqlite_master WHERE type = 'table'")
    }
    rows: dict[str, list[dict[str, Any]]] = {}
    for table in TABLES:
        if table.table not in present:
            rows[table.entity] = []
            continue
        order = ", ".join(_quote(c) for c in table.order_by)
        rows[table.entity] = [
            _decode(table, r) for r in brain.storage.fetchall(f"SELECT * FROM {_quote(table.table)} ORDER BY {order}")
        ]
    return rows


# -- the files ---------------------------------------------------------------------------------------
def _write_json(path: Path, payload: Any) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return path


def dataset_folder(out: Path, structure: dict[str, Any], parts: dict[str, list[dict[str, Any]]], source: str) -> Path:
    """The six files of ``structure`` in the Roy's Brain envelope, named as the owner's script names them
    (what Kletz's ``read_dataset_folder`` maps back to the entity fields), plus ``structure.json`` with
    the head the files cannot carry."""
    folder = out / STRUCTURES_DIR / str(structure["slug"])
    signals_of: dict[str, list[str]] = {}
    for s in parts["signals"]:
        signals_of.setdefault(str(s["edge"]), []).append(str(s["signal_id"]))
    files: dict[str, list[dict[str, Any]]] = {
        "controllers": [
            {"id": c["key"], "order": c["order"], "name": c["name"], "kind": c["kind"],
             "responsibility": c.get("responsibility") or "", "process_model_title": c.get("process_model_title") or "",
             "note": c.get("note") or ""}
            for c in parts["controllers"]
        ],
        "algorithms": [
            {"id": a["code"], "controller_id": a["controller"], "order": a["order"], "name": a.get("name") or ""}
            for a in parts["algorithms"]
        ],
        "process_model_entries": [
            {"id": f"{p['controller']}.pm.{p['order'] or index:02d}", "controller_id": p["controller"],
             "order": p["order"], "heading": p.get("heading") or "", "control_actions": p["control_actions"],
             "variables": p["variables"], "updated_by": p["updated_by"], "references": p["references"]}
            for index, p in enumerate(parts["process_model_entries"], start=1)
        ],
        "signals": [
            {"id": f"{s['signal_id']}@{s['edge']}", "signal_id": s["signal_id"], "kind": s.get("kind"),
             "name": s.get("name") or "", "edge_id": s["edge"], "order": s["order"],
             "placeholder": bool(s["placeholder"])}
            for s in parts["signals"]
        ],
        "edges": [
            {"id": e["key"], "order": e["order"], "kind": e["kind"], "source_id": e["source"], "target_id": e["target"],
             "signal_ids": signals_of.get(str(e["key"]), [])}
            for e in parts["edges"]
        ],
        "notes": [
            {"id": f"note.{n['order'] or index}", "order": n["order"], "text": n["text"],
             "attached_to": n.get("attached_to")}
            for index, n in enumerate(parts["notes"], start=1)
        ],
    }
    stamp = _timestamp()
    for name in DATASET_FILES:
        _write_json(folder / f"{name}.json", {
            "schema_version": 1, "exported_at": stamp, "source": source, "structure": structure["slug"],
            "entity": name, "row_count": len(files[name]), "rows": files[name],
        })
    _write_json(folder / "structure.json", structure)
    return folder


def _copy_bundle(source_bundle: Any, db_path: str, target: Path) -> Path | None:
    """The Hopper bundle an analysis was read from, copied under ``target`` when it is on this machine
    (as given, else relative to the database's folder); ``None`` when it is not."""
    if not source_bundle:
        return None
    candidates = [Path(str(source_bundle))]
    if not candidates[0].is_absolute() and db_path != ":memory:":
        candidates.append(Path(db_path).parent / str(source_bundle))
    for folder in candidates:
        if folder.is_dir() and (folder / "manifest.json").is_file():
            if target.exists():
                shutil.rmtree(target)
            shutil.copytree(folder, target)
            return target
    return None


def write_workbook(path: Path, rows: dict[str, list[dict[str, Any]]]) -> Path:
    """The thirteen sheets headed as the 0.10.1 workbook headed them (the field labels, ``uid`` last and
    hidden), lists as JSON text, booleans as 0 and 1; what ``asi-kletz import excel`` reads back."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for table in TABLES:
        ws = wb.create_sheet(table.sheet[:31])
        for col, header in enumerate(table.headers, start=1):
            ws.cell(row=1, column=col, value=header).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(len(table.headers))].hidden = True
        for r, record in enumerate(rows[table.entity], start=2):
            for col, column in enumerate(table.columns, start=1):
                value = record.get(column.name)
                if value is None or value == []:
                    continue
                if column.kind == "json":
                    value = json.dumps(list(value), ensure_ascii=False)
                elif column.kind == "bool":
                    value = 1 if value else 0
                ws.cell(row=r, column=col, value=value)
            ws.cell(row=r, column=len(table.headers), value=record.get("uid"))
        ws.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _quoted(path: Path | str) -> str:
    return f'"{path}"'


def write_handover(
    brain: Brain, out: str | Path | None = None, actor: str = "python"
) -> tuple[dict[str, Any], list[str]]:
    """Write everything under ``out`` (default ``handover/`` under the out folder); returns the summary
    (``out``, ``database``, ``counts``, ``files``, ``structures``, ``analyses``, ``kletz``) and the
    warnings (a bundle not on this machine). Raises ``CommandError`` when there is nothing to hand over."""
    rows = read_tables(brain)
    counts = {entity: len(rows[entity]) for entity in ENTITIES}
    if not any(counts.values()):
        raise CommandError(
            "nothing to hand over: the database holds no control structure and no STPA analysis "
            "(the thirteen tables of 0.9.0 to 0.10.1 are absent or empty)"
        )
    folder = Path(out) if out else Path(brain.out_dir) / DEFAULT_DIR
    folder.mkdir(parents=True, exist_ok=True)
    source = f"asi_dewey {__version__} handover of {brain.db_path}"
    warnings: list[str] = []
    kletz: list[str] = []

    payload = {
        "schema_version": SCHEMA_VERSION, "app_version": __version__, "exported_at": _timestamp(),
        "handover": {"from": brain.db_path, "for": f"asi_kletz {KLETZ_VERSION} load --mode merge", "tool": source},
        "entities": {entity: [{k: v for k, v in r.items() if k != "id"} for r in rows[entity]] for entity in ENTITIES},
    }
    json_path = _write_json(folder / JSON_FILE, payload)
    kletz.append(f"asi-kletz load {_quoted(json_path)} --mode merge")

    structures: list[dict[str, Any]] = []
    for head in rows["control_structure"]:
        slug = str(head["slug"])
        parts = {
            key: [{k: v for k, v in r.items() if k not in ("structure", "uid", "created_at", "updated_at")}
                  for r in rows[entity] if r.get("structure") == slug]
            for entity, key in _PARTS.items()
        }
        structure = {k: v for k, v in head.items() if k not in ("created_at", "updated_at")}
        dataset = dataset_folder(folder, structure, parts, source)
        command = (
            f"asi-kletz cs new from_datasets --from {_quoted(dataset)} --name {_quoted(head['name'])} --slug {slug}"
            + (f" --system {_quoted(head['system'])}" if head.get("system") else "")
        )
        kletz.append(command)
        structures.append({
            "slug": slug, "name": head["name"], "folder": str(dataset),
            "counts": {key: len(v) for key, v in parts.items()}, "kletz": command,
        })

    analyses: list[dict[str, Any]] = []
    for head in rows["stpa_analysis"]:
        slug = str(head["slug"])
        bundle = _copy_bundle(head.get("source_bundle"), brain.db_path, folder / BUNDLES_DIR / slug)
        entry: dict[str, Any] = {
            "slug": slug, "name": head["name"], "bundle": str(bundle) if bundle else None,
            "counts": {
                entity: sum(1 for r in rows[entity] if r.get("analysis") == slug) for entity in STPA_ENTITIES[1:]
            },
        }
        if bundle:
            entry["kletz"] = (
                f"asi-kletz stpa new from_bundle --from {_quoted(bundle)} --name {_quoted(head['name'])} --slug {slug}"
            )
            kletz.append(entry["kletz"])
        elif head.get("source_bundle"):
            warnings.append(
                f"analysis {slug}: its bundle {head['source_bundle']} is not on this machine; Kletz loads the rows "
                f"from {JSON_FILE} and cannot read the page again until the bundle is beside it"
            )
        analyses.append(entry)

    workbook = write_workbook(folder / WORKBOOK_FILE, rows)
    manifest = {
        "tool": source, "for": f"asi_kletz {KLETZ_VERSION}", "taken_at": payload["exported_at"],
        "database": brain.db_path, "counts": counts,
        "files": {"json": str(json_path), "workbook": str(workbook)},
        "structures": structures, "analyses": analyses, "kletz": kletz, "warnings": warnings,
    }
    manifest_path = _write_json(folder / MANIFEST_FILE, manifest)
    data = {
        "out": str(folder), "database": brain.db_path, "counts": counts,
        "files": {"json": str(json_path), "workbook": str(workbook), "manifest": str(manifest_path)},
        "structures": structures, "analyses": analyses, "kletz": kletz,
    }
    brain.storage.audit(actor, "handover", None, None, {"out": str(folder), **counts})
    brain.events.publish("handover.written", out=str(folder), counts=counts, actor=actor)
    return data, warnings

